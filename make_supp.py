# -*- coding: utf-8 -*-
import sys, pickle
sys.path.insert(0, ".")
import pandas as pd, docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from core import LEARNERS

R = "results/"
doc = docx.Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)

def table(title, header, rows, note=None):
    p = doc.add_paragraph(); r = p.add_run(title); r.bold = True
    t = doc.add_table(rows=1, cols=len(header)); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, header):
        c.text = ""; rr = c.paragraphs[0].add_run(h); rr.bold = True; rr.font.size = Pt(9)
        rr.font.color.rgb = RGBColor(255, 255, 255); shade(c, "2E75B6")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = ""; rr = c.paragraphs[0].add_run(str(v)); rr.font.size = Pt(9)
            if i % 2 == 0: shade(c, "EAF1F8")
    if note:
        n = doc.add_paragraph(); rn = n.add_run(note); rn.italic = True; rn.font.size = Pt(8.5)
    doc.add_paragraph()

h = doc.add_paragraph(); r = h.add_run("Supplementary Material"); r.bold = True; r.font.size = Pt(14)
doc.add_paragraph("Sensor Features Versus the Wear–Time Trend: A Leakage-Free Cross-Tool Evaluation of Tool-Wear "
                  "Prediction Under Limited Labels on the PHM2010 Milling Benchmark")

# S1a feature definitions
table("Table S1a. Feature definitions (computed per cut on the steady-state segment, first and last 10% of samples removed).",
      ["Family", "Channels", "Features", "Settings"],
      [["Time domain", "all 7", "mean, std, RMS, skewness, kurtosis, peak-to-peak, crest factor", "—"],
       ["Kinematic spectral", "force x/y/z, vib x/y/z",
        "E_fs; E_ft1…E_ft5; ratio E_fs/E_ft1; E_sideband (f_t ± f_s); spectral centroid",
        "Welch PSD, nperseg 16,384 (≈3.05 Hz); ±5 Hz bands; f_s = 173.33 Hz, f_t = 520 Hz; centroid 0–10 kHz; "
        "energies and ratio log10-transformed"],
       ["Wavelet packet", "force x/y/z, vib x/y/z", "relative energies wp0…wp7", "db4, level 3, frequency-ordered nodes"],
       ["AE low band", "AE-RMS", "E_low", "Welch PSD energy below 1 kHz, log10"],
       ["1D-CNN input", "all 7", "raw window", "decimated to 10 kHz (zero-phase), central 4,096 samples"]],
      "Total: 6 × 24 + 8 = 152 features per cut.")

# S1b grids
gr = []
for name, (_, g, stoch) in LEARNERS.items():
    keys = list(g[0].keys())
    vals = {k: sorted({str(x[k]) for x in g}, key=lambda v: g and [str(x[k]) for x in g].index(v)) for k in keys}
    gr.append([name, "; ".join(f"{k} ∈ {{{', '.join(vals[k])}}}" for k in keys), "0–4 (averaged)" if stoch else "—"])
gr.append(["Fixed settings", "Ridge proxy for selection: α = 1, up to 20 features; Random Forest n_estimators = 300; "
           "Gradient Boosting / XGBoost subsample = 0.8; MLP max_iter = 3,000; inputs and target standardised", ""])
gr.append(["1D-CNN", "epochs ∈ {5, 10, …, 150} chosen on inner split; Adam lr 1e-3, weight decay 1e-4, batch 8, "
           "dropout 0.3, 2,048-sample random crops, 5 test crops", "0–4 (averaged)"])
table("Table S1b. Hyperparameter grids (grouped inner split: fit on one training cutter, validate on the other).",
      ["Learner", "Grid", "Seeds"], gr)

# S1c selected configs
meta = pickle.load(open(R + "folds_main_max.pkl", "rb"))
cnn = pd.read_pickle(R + "cnn_meta_max_k20.pkl")
rows = []
for o, c in zip(meta, cnn):
    rows.append([o["test"], o["baseline_variant"], str(len(o["selected"]))] +
                ["; ".join(f"{k}={v}" for k, v in o["configs"][m].items()) for m in
                 ["Ridge", "Lasso", "SVR", "Random Forest", "Gradient Boosting", "XGBoost", "MLP"]] + [str(c["epochs"])])
rowsT = list(map(list, zip(*rows)))
labels = ["Baseline variant", "No. selected features", "Ridge", "Lasso", "SVR", "Random Forest", "Gradient Boosting",
          "XGBoost", "MLP", "1D-CNN epochs"]
table("Table S1c. Configuration selected in each outer fold (target VB_max, k = 20).",
      ["Setting", "held-out c1", "held-out c4", "held-out c6"], [[l] + r for l, r in zip(labels, rowsT[1:])])

table("Table S1d. Software versions.", ["Stage", "Versions"],
      [["Feature extraction (author's computer)", "NumPy 2.2.6, pandas 2.3.3, SciPy 1.15.3, PyWavelets 1.8.0"],
       ["Modelling and statistics", "Python 3.13, NumPy 2.5.3, SciPy 1.18.1, scikit-learn 1.9.1, XGBoost 3.4.1, "
        "PyTorch 2.14.1, SHAP 0.54.1, Matplotlib 3.11.2"]])

# S2 mean target
S = pd.read_pickle(R + "summary_main_mean.pkl")
f = lambda x: f"{x:+.4f}".replace("-", "−")
rows = [[r.model, f"{r.RMSE_pooled:.4f}", f"{r.R2_pooled:.3f}",
         "—" if r.model == "Cut-index baseline" else f"{f(r.dRMSE)} [{f(r.dRMSE_CI[0])}, {f(r.dRMSE_CI[1])}]",
         f"{r.RMSE_c1:.4f} / {r.RMSE_c4:.4f} / {r.RMSE_c6:.4f}"] for r in S.itertuples()]
table("Table S2. Sensitivity analysis: target = mean flank wear of the three flutes (VB_mean), k = 20.",
      ["Model", "RMSE (mm), pooled", "R², pooled", "ΔRMSE vs baseline (mm) [95% CI]", "RMSE c1 / c4 / c6 (mm)"], rows,
      "Same nested pipeline as Table 2 of the main text; feature selection and tuning repeated for this target.")

# S3 threshold sensitivity
X = pd.read_csv(R + "crossing_max.csv")
rows = []
for lim, g in X.groupby("lim"):
    r = [f"{lim:.2f}"]
    for m in ["Cut-index baseline", "Random Forest"]:
        h = g[g.model == m].set_index("cutter")
        r.append(" | ".join(f"{c}: {int(h.loc[c,'early'])}/{int(h.loc[c,'late'])}" for c in ["c1", "c4", "c6"]))
        r.append(str(int(h.early.sum() + h.late.sum())))
    rows.append(r)
table("Table S3. Threshold-crossing errors (Δearly/Δlate, cuts) for VB_lim = 0.10–0.16 mm.",
      ["VB_lim (mm)", "Baseline per cutter", "Baseline total", "Random Forest per cutter", "Random Forest total"], rows)


# S4 threshold sensitivity, normalised features
X = pd.read_csv(R + "crossing_norm5.csv")
rows = []
for lim, g in X.groupby("lim"):
    r = [f"{lim:.2f}"]
    for m in ["Cut-index baseline", "Ridge + cut index", "MLP", "Random Forest"]:
        h = g[g.model == m].set_index("cutter")
        r.append(" | ".join(f"{c}: {int(h.loc[c,'early'])}/{int(h.loc[c,'late'])}" for c in ["c1", "c4", "c6"]) +
                 f"  (Σ {int(h.early.sum() + h.late.sum())})")
    rows.append(r)
table("Table S4. Post-hoc analysis: threshold-crossing errors (Δearly/Δlate, cuts; Σ = total) with per-tool "
      "normalised features (n_cal = 5).",
      ["VB_lim (mm)", "Cut-index baseline", "Ridge + cut index", "MLP", "Random Forest"], rows)

# S5 NASA all models
NR = "nasa/results/"
A_ = {k: pd.read_pickle(NR + f"summary_{k}.pkl").set_index("model") for k in ["setting_raw", "setting_norm", "case_raw"]}
g3 = lambda t: f"[{t[0]:+.3f}, {t[1]:+.3f}]".replace("-", "−")
rows = []
for m in A_["setting_raw"].index:
    r = [m]
    for k in ["setting_raw", "setting_norm", "case_raw"]:
        x = A_[k].loc[m]
        r.append(f"{x.RMSE:.3f}" + ("" if m == "Process baseline" else f" {g3(x.d_proc_CI)}"))
    rows.append(r)
table("Table S5. NASA Ames milling data: held-out RMSE (mm) and ΔRMSE vs process baseline [95% cluster-bootstrap CI] "
      "for all models.", ["Model", "Leave-one-setting-out, raw", "Leave-one-setting-out, normalised",
                          "Leave-one-case-out, raw"], rows,
      "Sensor models without the suffix use selected sensor features plus run time, depth of cut, feed and material.")
doc.save("Supplementary_Material.docx"); print("ok")
